#!/usr/bin/env python3
"""把 book.json 渲染成排版规整的 Word 文档。

用法:
  PYTHONPATH=pylibs python3 build_docx.py
"""
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from convert import split_bold

BASE = os.path.dirname(os.path.abspath(__file__))
os.chdir(BASE)
OUT = '如何明白圣经.docx'

BODY_CN = '宋体'
HEAD_CN = '黑体'
BODY_EN = 'Times New Roman'
HEAD_EN = 'Arial'


def set_font(run, cn=BODY_CN, en=BODY_EN, size=12, bold=False, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def style_run(style, cn, en, size, bold=False):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def clean(text):
    """去掉段落前的全角/半角缩进空格，缩进交给 Word 样式控制。"""
    return text.replace('\u3000', ' ').strip()


def write(p, text, **font):
    """按 **粗体** 拆 run 写入段落。"""
    for seg, bold in split_bold(text):
        set_font(p.add_run(seg), bold=bold, **font)


def add_page_numbers(section):
    """页脚居中显示「第 N 页」，用 PAGE 域实现。"""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run('第 '), size=10, color=RGBColor(0x70, 0x70, 0x70))
    run = p.add_run()
    set_font(run, size=10, color=RGBColor(0x70, 0x70, 0x70))
    fld = run._element.makeelement(qn('w:fldSimple'), {qn('w:instr'): 'PAGE'})
    run._element.addnext(fld)
    set_font(p.add_run(' 页'), size=10, color=RGBColor(0x70, 0x70, 0x70))


def first_line_indent(p, chars=200):
    """按东亚字符数设置首行缩进（200 = 2 字符）。"""
    pf = p._p.get_or_add_pPr()
    ind = pf.find(qn('w:ind'))
    if ind is None:
        ind = pf.makeelement(qn('w:ind'), {})
        pf.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), str(int(chars / 100 * 12 * 20)))  # 后备值，单位 twip


def main():
    data = json.load(open('book.json', encoding='utf-8'))
    meta, chapters = data['meta'], data['chapters']

    doc = Document()

    # 页面：A4，适中页边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.8)

    # 正文默认样式
    normal = doc.styles['Normal']
    style_run(normal, BODY_CN, BODY_EN, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    add_page_numbers(sec)

    # 封面
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta['book']), cn=HEAD_CN, en=HEAD_EN, size=30, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta['title_en']), size=15, color=RGBColor(0x55, 0x55, 0x55))
    doc.add_paragraph()
    for label, value in (('作者', meta['author']), ('翻译', meta['translator']),
                         ('英文版', meta['publisher_en']), ('来源', meta['source'])):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f'{label}：{value}'), size=11)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 目录
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run('目　录'), cn=HEAD_CN, en=HEAD_EN, size=20, bold=True)

    cur = None
    for ch in chapters:
        if ch['part'] and ch['part'] != cur:
            cur = ch['part']
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            set_font(p.add_run(cur), cn=HEAD_CN, en=HEAD_EN, size=13, bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.8)
        set_font(p.add_run(ch['title']), size=12)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 正文
    cur = None
    for idx, ch in enumerate(chapters):
        if ch['part'] and ch['part'] != cur:
            cur = ch['part']
            if idx:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(60)
            set_font(p.add_run(cur), cn=HEAD_CN, en=HEAD_EN, size=22, bold=True)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        for b in ch['blocks']:
            t, text = b['type'], b['text']
            if t == 'h1':
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(20)
                write(p, clean(text), cn=HEAD_CN, en=HEAD_EN, size=20)
            elif t == 'h2':
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                write(p, clean(text), cn=HEAD_CN, en=HEAD_EN, size=15)
            elif t in ('h3', 'h4', 'h5', 'h6'):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                write(p, clean(text), cn=HEAD_CN, en=HEAD_EN, size=13)
            elif t == 'quote':
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(4)
                write(p, clean(text), size=10, color=RGBColor(0x60, 0x60, 0x60))
            elif t == 'li':
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(3)
                write(p, clean(text), size=12)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                first_line_indent(p, 200)
                write(p, clean(text), size=12)

        if idx < len(chapters) - 1 and chapters[idx + 1]['part'] == ch['part']:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.save(OUT)
    print(f'saved {OUT}  chapters={len(chapters)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
