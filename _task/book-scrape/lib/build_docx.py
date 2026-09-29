#!/usr/bin/env python3
"""把 book.json 渲染成排版规整的 Word 文档。

用法:
  python3 lib/build_docx.py <书名目录>
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bookkit

bookkit.use_pylibs()

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from convert import split_bold  # noqa: E402

PAGE_SIZES = {'A4': (21.0, 29.7), 'A5': (14.8, 21.0), 'Letter': (21.59, 27.94)}


def set_font(run, cn, en, size, bold=False, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def style_run(style, cn, en, size, bold=False):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def clean(text):
    """去掉段落前的全角/半角缩进空格，缩进交给 Word 样式控制。"""
    return text.replace('\u3000', ' ').strip()


class Renderer:
    def __init__(self, cfg, doc):
        self.cfg = cfg
        self.doc = doc
        self.body_cn = cfg['body_cn']
        self.head_cn = cfg['head_cn']
        self.body_en = cfg['body_en']
        self.head_en = cfg['head_en']
        self.body_size = cfg['body_size']

    def body(self, p, text, **font):
        font.setdefault('cn', self.body_cn)
        font.setdefault('en', self.body_en)
        font.setdefault('size', self.body_size)
        self._runs(p, text, font)

    def head(self, p, text, size, **font):
        font.setdefault('cn', self.head_cn)
        font.setdefault('en', self.head_en)
        font.setdefault('size', size)
        font.setdefault('bold', True)
        self._runs(p, text, font)

    @staticmethod
    def _runs(p, text, font):
        bold_default = font.pop('bold', False)
        for seg, bold in split_bold(clean(text)):
            set_font(p.add_run(seg), bold=bold or bold_default, **font)

    def page_numbers(self, section):
        if not self.cfg.get('page_number'):
            return
        gray = RGBColor(0x70, 0x70, 0x70)
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run('第 '), self.body_cn, self.body_en, 10, color=gray)
        run = p.add_run()
        set_font(run, self.body_cn, self.body_en, 10, color=gray)
        fld = run._element.makeelement(qn('w:fldSimple'), {qn('w:instr'): 'PAGE'})
        run._element.addnext(fld)
        set_font(p.add_run(' 页'), self.body_cn, self.body_en, 10, color=gray)

    @staticmethod
    def first_line_indent(p, chars=200):
        """按东亚字符数设置首行缩进（200 = 2 字符）。"""
        pf = p._p.get_or_add_pPr()
        ind = pf.find(qn('w:ind'))
        if ind is None:
            ind = pf.makeelement(qn('w:ind'), {})
            pf.append(ind)
        ind.set(qn('w:firstLineChars'), str(chars))
        ind.set(qn('w:firstLine'), str(int(chars / 100 * 12 * 20)))  # 后备值，单位 twip


def main():
    bookkit.book_dir(sys.argv[1] if len(sys.argv) > 1 else None)
    data = json.load(open('book.json', encoding='utf-8'))
    meta, chapters, cfg = data['meta'], data['chapters'], data['docx']

    doc = Document()
    r = Renderer(cfg, doc)

    sec = doc.sections[0]
    w, h = PAGE_SIZES.get(cfg.get('page', 'A4'), PAGE_SIZES['A4'])
    sec.page_width, sec.page_height = Cm(w), Cm(h)
    vertical, horizontal = cfg['margin_cm']
    sec.top_margin = sec.bottom_margin = Cm(vertical)
    sec.left_margin = sec.right_margin = Cm(horizontal)

    normal = doc.styles['Normal']
    style_run(normal, cfg['body_cn'], cfg['body_en'], cfg['body_size'])
    normal.paragraph_format.line_spacing = cfg['line_spacing']
    normal.paragraph_format.space_after = Pt(6)

    r.page_numbers(sec)

    # 封面
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r.head(p, meta['book'], 30)
    if meta.get('title_en'):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r.body(p, meta['title_en'], size=15, color=RGBColor(0x55, 0x55, 0x55))
    doc.add_paragraph()
    cover = [('作者', meta.get('author')), ('翻译', meta.get('translator')),
             ('英文版', meta.get('publisher_en')), ('来源', meta.get('source'))]
    for label, value in cover:
        if not value:
            continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r.body(p, f'{label}：{value}', size=11)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 目录
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r.head(p, '目　录', 20)

    cur = None
    for ch in chapters:
        if ch.get('part') and ch['part'] != cur:
            cur = ch['part']
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            r.head(p, cur, 13)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.8)
        r.body(p, ch['title'])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 正文
    cur = None
    for idx, ch in enumerate(chapters):
        if ch.get('part') and ch['part'] != cur:
            cur = ch['part']
            if idx:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(60)
            r.head(p, cur, 22)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        for b in ch['blocks']:
            t, text = b['type'], b['text']
            if t == 'h1':
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(20)
                r.head(p, text, 20)
            elif t == 'h2':
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                r.head(p, text, 15)
            elif t in ('h3', 'h4', 'h5', 'h6'):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                r.head(p, text, 13)
            elif t == 'quote':
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(4)
                r.body(p, text, size=10, color=RGBColor(0x60, 0x60, 0x60))
            elif t == 'li':
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.line_spacing = cfg['line_spacing']
                p.paragraph_format.space_after = Pt(3)
                r.body(p, text)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r.first_line_indent(p, 200)
                r.body(p, text)

        if idx < len(chapters) - 1 and chapters[idx + 1].get('part') == ch.get('part'):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out = cfg.get('filename') or f"{meta['book']}.docx"
    if not out.endswith('.docx'):
        out += '.docx'
    doc.save(out)
    print(f'saved {out}  chapters={len(chapters)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
