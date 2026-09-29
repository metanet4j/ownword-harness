import json, os, re, glob, sys
sys.path.insert(0, os.path.join(os.getcwd(), 'pylibs'))
from bs4 import BeautifulSoup
from docx import Document

d = sys.argv[1]
mf = json.load(open(f'{d}/manifest.json', encoding='utf-8'))
doc = Document(f'{d}/' + (mf['docx'].get('filename') or f"{mf['book']}.docx"))
alltext = ''.join(p.text for p in doc.paragraphs)
dn = re.sub(r'\s+', '', alltext)
bad = []
for ch in mf['chapters']:
    s = BeautifulSoup(open(f"{d}/raw/{ch['slug']}.html", encoding='utf-8').read(), 'lxml')
    c = s.select_one('div.entry-content')
    for t in c.select('audio,script,style'): t.decompose()
    ps = [re.sub(r'\s+','',p.get_text()) for p in c.find_all('p') if len(re.sub(r'\s+','',p.get_text()))>30]
    md = open(f"{d}/{ch['slug']}.md", encoding='utf-8').read()
    md = re.sub(r'^#+.*$', '', md, flags=re.M)      # 标题行
    md = re.sub(r'^>.*$', '', md, flags=re.M)       # 音频引用行
    mdn = re.sub(r'\s+', '', md.replace('**', '').replace('\\*', '*'))
    miss_docx = [p[:30] for p in ps if p not in dn]
    miss_md = [p[:30] for p in ps if p not in mdn]
    if miss_docx or miss_md:
        bad.append((ch['slug'], len(miss_md), len(miss_docx), (miss_md or miss_docx)[0]))
print('docx 段落:', len(doc.paragraphs), '| 字符:', len(dn), '| 星号:', alltext.count('*'))
print('逐章核对:', bad if bad else '无缺失（Markdown 与 Word 全部命中）')
links = re.findall(r'\]\((.+?\.md)\)', open(f'{d}/目录.md', encoding='utf-8').read())
print('目录链接:', len(links), '| 失效:', [l for l in links if not os.path.exists(f'{d}/{l}')] or '无')
print('分章文件:', len([f for f in glob.glob(f'{d}/*.md') if os.path.basename(f) != '目录.md']), '篇 + 目录.md')
